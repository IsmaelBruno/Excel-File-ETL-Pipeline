"""Lecture Excel : feuille, en-tête, lignes brutes liées aux colonnes."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from atg_excel_pipeline.config import Settings
from atg_excel_pipeline.headers import (
    StructureError,
    bind_columns,
    find_header_row,
    normalize_header,
)
from atg_excel_pipeline.schemas import Flow, FlowSpec, spec_for


@dataclass(frozen=True)
class RawRow:
    source_row: int
    values: tuple[object, ...]
    bound: dict[str, object]


@dataclass(frozen=True)
class ParsedWorkbook:
    path: Path
    flow: Flow
    spec: FlowSpec
    sheet_name: str
    header_row: int
    source_headers: tuple[object, ...]
    rows: tuple[RawRow, ...]
    warnings: tuple[str, ...]


def _select_sheet(
    sheetnames: list[str],
    spec: FlowSpec,
    *,
    allow_fallback: bool,
) -> str:
    # Correspondance exacte en priorité
    if spec.sheet_name in sheetnames:
        return spec.sheet_name

    # Flexibilité uniquement pour Entrée et Passage
    flexible_sheets = {
        "RDR56REBRE_FR",
        "RDR56REBSL_FR",
    }

    def normalize_sheet(name: str) -> str:
        return "".join(
            char.upper()
            for char in name
            if char.isalnum()
        )

    if spec.sheet_name in flexible_sheets:
        target = normalize_sheet(spec.sheet_name)

        candidates = [
            name for name in sheetnames
            if normalize_sheet(name).startswith(target)
        ]

        if len(candidates) == 1:
            return candidates[0]

        if len(candidates) > 1:
            raise StructureError(
                f"Plusieurs feuilles correspondent à "
                f"{spec.sheet_name!r} : {candidates}"
            )

    # Conserver le comportement existant des autres flux
    if allow_fallback and sheetnames:
        return sheetnames[0]

    raise StructureError(
        f"Feuille {spec.sheet_name!r} absente "
        f"(feuilles disponibles : {sheetnames})."
    )


def parse_xlsx(path: Path, spec: FlowSpec, settings: Settings) -> ParsedWorkbook:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet_name = _select_sheet(
            workbook.sheetnames,
            spec,
            allow_fallback=settings.allow_first_sheet_fallback,
        )
        worksheet = workbook[sheet_name]
        buffer: list[tuple[object, ...]] = []
        extra: list[tuple[int, tuple[object, ...]]] = []
        header_found_at: int | None = None
        markers = {normalize_header(m) for m in spec.header_markers}

        for index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            values = tuple(row)
            if header_found_at is None:
                buffer.append(values)
                first = values[0] if values else None
                if normalize_header(first) in markers:
                    header_found_at = index
                    continue
                if index >= settings.header_scan_max_rows:
                    break
                continue
            extra.append((index, values))

        first_column = [row[0] if row else None for row in buffer]
        header_row = find_header_row(
            first_column,
            spec,
            max_rows=settings.header_scan_max_rows,
            allow_fallback=settings.allow_header_fallback,
        )
        if header_row > len(buffer):
            raise StructureError(f"Ligne d'en-tête {header_row} hors scan.")
        headers = list(buffer[header_row - 1])
        binding = bind_columns(headers, spec)
        warnings = list(binding.warnings)

        raw_rows: list[RawRow] = []
        for offset, values in enumerate(buffer[header_row:], start=header_row + 1):
            raw_rows.append(_bind_row(offset, values, binding.bound))
        for index, values in extra:
            raw_rows.append(_bind_row(index, values, binding.bound))

        return ParsedWorkbook(
            path=path,
            flow=spec.flow,
            spec=spec,
            sheet_name=sheet_name,
            header_row=header_row,
            source_headers=tuple(headers),
            rows=tuple(raw_rows),
            warnings=tuple(warnings),
        )
    finally:
        workbook.close()


def parse_xlsb(path: Path, spec: FlowSpec, settings: Settings) -> ParsedWorkbook:
    from pyxlsb import open_workbook

    with open_workbook(str(path)) as workbook:
        sheet_name = _select_sheet(
            list(workbook.sheets),
            spec,
            allow_fallback=settings.allow_first_sheet_fallback,
        )
        buffer: list[tuple[object, ...]] = []
        extra: list[tuple[int, tuple[object, ...]]] = []
        header_found_at: int | None = None
        markers = {normalize_header(m) for m in spec.header_markers}
        with workbook.get_sheet(sheet_name) as sheet:
            for index, row in enumerate(sheet.rows(), start=1):
                values = tuple(cell.v for cell in row)
                if header_found_at is None:
                    buffer.append(values)
                    first = values[0] if values else None
                    if normalize_header(first) in markers:
                        header_found_at = index
                        continue
                    if index >= settings.header_scan_max_rows:
                        break
                    continue
                extra.append((index, values))

        first_column = [row[0] if row else None for row in buffer]
        header_row = find_header_row(
            first_column,
            spec,
            max_rows=settings.header_scan_max_rows,
            allow_fallback=settings.allow_header_fallback,
        )
        headers = list(buffer[header_row - 1])
        binding = bind_columns(headers, spec)
        raw_rows: list[RawRow] = []
        for offset, values in enumerate(buffer[header_row:], start=header_row + 1):
            raw_rows.append(_bind_row(offset, values, binding.bound))
        for index, values in extra:
            raw_rows.append(_bind_row(index, values, binding.bound))
        return ParsedWorkbook(
            path=path,
            flow=spec.flow,
            spec=spec,
            sheet_name=sheet_name,
            header_row=header_row,
            source_headers=tuple(headers),
            rows=tuple(raw_rows),
            warnings=tuple(binding.warnings),
        )


def parse_workbook(path: Path, flow: Flow, settings: Settings) -> ParsedWorkbook:
    spec = spec_for(flow)
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return parse_xlsx(path, spec, settings)
    if suffix == ".xlsb":
        return parse_xlsb(path, spec, settings)
    raise StructureError(f"Extension non prise en charge: {path.suffix}")


def _bind_row(source_row: int, values: tuple[object, ...], bound) -> RawRow:
    mapped: dict[str, object] = {}
    for column in bound:
        mapped[column.spec.field] = (
            values[column.source_index] if column.source_index < len(values) else None
        )
    return RawRow(source_row=source_row, values=values, bound=mapped)
