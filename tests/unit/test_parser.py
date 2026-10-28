from __future__ import annotations

from datetime import date, time
from pathlib import Path

from openpyxl import Workbook

from atg_excel_pipeline.config import Settings
from atg_excel_pipeline.parser import parse_workbook
from atg_excel_pipeline.schemas import ENTREE_COLUMNS, SORTIE_COLUMNS, Flow
from atg_excel_pipeline.transformer import transform_workbook


def _headers(columns) -> list[object]:
    width = max(c.position for c in columns)
    values: list[object] = [None] * width
    for column in columns:
        values[column.position - 1] = column.source_header
    return values


def _settings(tmp_path: Path) -> Settings:
    return Settings.from_environ(
        {
            "ENTREE_DIR": str(tmp_path / "in"),
            "SORTIE_DIR": str(tmp_path / "out"),
            "STAGING_DIR": str(tmp_path / "stg"),
        }
    )


def test_parse_et_transforme_entree(tmp_path: Path) -> None:
    path = tmp_path / "Entrées Sep-26 JdE.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RDR56REBRE_FR"
    for col, header in enumerate(_headers(ENTREE_COLUMNS), start=1):
        sheet.cell(15, col, header)
    sheet.cell(16, 1, "DEPOT")
    sheet.cell(16, 2, "GO")
    sheet.cell(16, 4, 12345.0)
    sheet.cell(16, 9, date(2026, 9, 1))
    sheet.cell(16, 13, 12.5)
    sheet.cell(16, 30, "ABESPTOTAL")
    sheet.cell(16, 31, time(12, 0, 47))
    sheet.cell(17, 1, None)
    sheet.cell(18, 1, "DEPOT")
    sheet.cell(18, 9, "pas-une-date")
    workbook.save(path)
    workbook.close()

    parsed = parse_workbook(path, Flow.ENTREE, _settings(tmp_path))
    assert parsed.header_row == 15
    transformed = transform_workbook(parsed, _settings(tmp_path))
    assert transformed.quality.valid_rows == 1
    assert transformed.quality.rejected_rows == 1
    assert transformed.quality.rows_blank_skipped >= 1
    row = transformed.valid_rows[0]
    assert row["n_bon"] == "12345"
    assert row["ege"] == 12.5
    assert row["date_principale"] == date(2026, 9, 1)
    assert row["date_c30"] == "ABESPTOTAL"
    assert transformed.quality.conversion_errors >= 1
    assert row["heure"] == time(12, 0, 47)


def test_parse_sortie_heure_et_produit(tmp_path: Path) -> None:
    path = tmp_path / "Passage Sep-26 JdE.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "RDR56REBSL_FR"
    for col, header in enumerate(_headers(SORTIE_COLUMNS), start=1):
        sheet.cell(14, col, header)
    sheet.cell(15, 1, "DEPOT")
    sheet.cell(15, 2, "GO")
    sheet.cell(15, 5, 10.4)
    sheet.cell(15, 7, date(2026, 9, 2))
    sheet.cell(15, 13, "1 234,5")
    sheet.cell(15, 30, 93005)
    sheet.cell(16, 1, "DEPOT")
    sheet.cell(16, 2, None)
    sheet.cell(16, 7, date(2026, 9, 3))
    workbook.save(path)
    workbook.close()

    parsed = parse_workbook(path, Flow.SORTIE, _settings(tmp_path))
    transformed = transform_workbook(parsed, _settings(tmp_path))
    assert transformed.quality.valid_rows == 1
    assert "PRODUIT_ABSENT" in transformed.quality.reject_motifs
    row = transformed.valid_rows[0]
    assert row["egs"] == 1234.5
    assert row["n_bsl"] == 10
    assert row["heure"] == time(9, 30, 5)
